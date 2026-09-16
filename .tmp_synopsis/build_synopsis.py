from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

OUT = r'D:\Frontend\obrix-build-obrix-connected-frontend\OBRIX_Frontend_Project_Synopsis.docx'

def font(run, size=12, bold=False, italic=False):
    run.font.name = 'Times New Roman'
    run._element.rPr.rFonts.set(qn('w:ascii'), 'Times New Roman')
    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Times New Roman')
    run.font.size = Pt(size); run.bold = bold; run.italic = italic

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

def borders(table):
    tblPr = table._tbl.tblPr; el = OxmlElement('w:tblBorders')
    for side in ('top','left','bottom','right','insideH','insideV'):
        node = OxmlElement('w:' + side); node.set(qn('w:val'),'single'); node.set(qn('w:sz'),'4'); node.set(qn('w:color'),'D9D9D9'); el.append(node)
    tblPr.append(el)

def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    node = OxmlElement('w:tblHeader')
    node.set(qn('w:val'), 'true')
    trPr.append(node)

doc=Document()
sec=doc.sections[0]
sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)

normal=doc.styles['Normal']; normal.font.name='Times New Roman'; normal._element.rPr.rFonts.set(qn('w:ascii'),'Times New Roman'); normal.font.size=Pt(12)
normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.15
for name, size in [('Title',24),('Heading 1',16),('Heading 2',13)]:
    st=doc.styles[name]; st.font.name='Times New Roman'; st._element.rPr.rFonts.set(qn('w:ascii'),'Times New Roman'); st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.space_before=Pt(12); st.paragraph_format.space_after=Pt(8)

def p(text='', align=None, size=12, bold=False, italic=False, style=None, before=None, after=None):
    para=doc.add_paragraph(style=style)
    if align is not None: para.alignment=align
    run=para.add_run(text); font(run,size,bold,italic)
    if before is not None: para.paragraph_format.space_before=Pt(before)
    if after is not None: para.paragraph_format.space_after=Pt(after)
    return para

def bullet(text):
    para=doc.add_paragraph(style='List Bullet'); para.paragraph_format.space_after=Pt(3); font(para.add_run(text),12); return para

def head(text, level=1):
    return p(text, style='Heading 1' if level==1 else 'Heading 2')

def page(): doc.add_page_break()

# Cover
p('OBRIX', WD_ALIGN_PARAGRAPH.CENTER, 24, True, before=90, after=20)
p('A', WD_ALIGN_PARAGRAPH.CENTER, 14, True)
p('SYNOPSIS REPORT OF PROJECT', WD_ALIGN_PARAGRAPH.CENTER, 14, True, after=14)
p('OBRIX: A CONSTRUCTION MARKETPLACE AND CONNECTED WORKSPACE', WD_ALIGN_PARAGRAPH.CENTER, 12, after=12)
p('For the course FRONT END WEB DEVELOPMENT (26CA205PCB)', WD_ALIGN_PARAGRAPH.CENTER, 12)
p('In partial fulfillment of requirements for the degree of', WD_ALIGN_PARAGRAPH.CENTER, 12)
p('MASTER OF COMPUTER APPLICATIONS (MCA)', WD_ALIGN_PARAGRAPH.CENTER, 16, True, after=34)
p('SUBMITTED BY:', WD_ALIGN_PARAGRAPH.CENTER, 12, True)
p('[TEAM MEMBER NAME(S) AND ROLL NUMBER(S)]', WD_ALIGN_PARAGRAPH.CENTER, 14, after=30)
p('[MONTH YEAR]', WD_ALIGN_PARAGRAPH.CENTER, 12, after=45)
p('SUBMITTED TO:', WD_ALIGN_PARAGRAPH.CENTER, 12, True)
p('[FACULTY NAME]', WD_ALIGN_PARAGRAPH.CENTER, 12)
p('[DEPARTMENT NAME]', WD_ALIGN_PARAGRAPH.CENTER, 12)
p('[INSTITUTE / UNIVERSITY NAME]', WD_ALIGN_PARAGRAPH.CENTER, 12)
p('[INSTITUTE ADDRESS]', WD_ALIGN_PARAGRAPH.CENTER, 12)

page(); head('TABLE OF CONTENTS')
sections=['Abstract','Introduction and Background','Problem Statement','Objectives','Scope of the Project','Existing and Proposed System','Functional Requirements','Non Functional Requirements','System Modules','Methodology','Technology Stack','Hardware and Software Requirements','System Design and Architecture','Data and State Management','Testing Strategy','Project Timeline','Expected Outcomes','Limitations and Future Scope','Conclusion','References']
t=doc.add_table(rows=1, cols=3); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'; borders(t)
for i,txt in enumerate(['No.','Section','Page']):
    c=t.rows[0].cells[i]; c.text=txt; shade(c,'D9EAD3'); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for r in c.paragraphs[0].runs: font(r,11,True)
for i,s in enumerate(sections,1):
    cells=t.add_row().cells; cells[0].text=str(i); cells[1].text=s; cells[2].text='[Auto]'
    for c in cells:
        c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for para in c.paragraphs:
            for r in para.runs: font(r,11)

content = [
('ABSTRACT', ['OBRIX is a responsive frontend web application that demonstrates a connected construction marketplace for companies, contractors, and skilled workers. It brings project discovery, bidding, workforce coordination, attendance, payment recording, expense tracking, equipment rental requests, and notifications into one browser-based interface.', 'The application is implemented as a React single-page application with Vite. It uses seeded mock data and browser local storage to simulate a shared workspace. The prototype is intentionally a frontend demonstration: no real money, user verification, database, server API, or live equipment marketplace is connected.']),
('INTRODUCTION AND BACKGROUND', ['Construction projects depend on timely coordination among project owners, contractors, workers, and equipment providers. In fragmented workflows, project information, bids, attendance records, payments, and equipment requests are often handled through separate channels. This makes progress difficult to trace and creates avoidable delays.', 'OBRIX presents one unified digital interface for these construction stakeholders. Public marketplace screens help visitors explore projects, equipment, contractors, and workers. Role-specific workspaces then expose the information and actions relevant to a company, contractor, or worker.']),
('PROBLEM STATEMENT', ['Construction participants need a clear and role-appropriate way to discover opportunities and keep related project activity visible. A company needs to publish work and compare proposals; a contractor needs to bid, form a site team, record daily attendance, and track expenses; a worker needs visibility into assigned work, attendance, earnings, and payment history. This project prototypes those connected journeys in one responsive frontend application.']),
('OBJECTIVES', ['The project objectives are:'])]

for title, pars in content:
    page(); head(title)
    for x in pars: p(x)
    if title=='OBJECTIVES':
        for x in ['Create a responsive marketplace interface for construction projects, equipment, contractors, and workers.','Provide three role-based demo workspaces for company, contractor, and worker users.','Support frontend workflows for project publication, bidding, bid approval, worker assignment, attendance, payments, expenses, equipment rentals, and notifications.','Use realistic seeded data and local persistence to demonstrate interconnected updates without a backend service.','Apply accessible interaction patterns such as semantic controls, keyboard-focusable dialogs, labels, responsive navigation, and a skip link.']: bullet(x)

page(); head('SCOPE OF THE PROJECT')
p('The current scope is a frontend prototype and demonstration environment. The following features are implemented in the repository:')
for x in ['Public landing page, marketplace listings, filters, sort controls, project details, equipment details, and demo authentication screens.','Company workflows for creating and editing open projects, reviewing bids, accepting or rejecting bids, tracking payments, and managing equipment rental requests.','Contractor workflows for browsing projects, submitting bids, monitoring active projects, assigning workers, recording attendance, recording expenses, and settling simulated wages.','Worker workflows for viewing assignments, attendance, earnings, payment history, skills, availability, reviews, and profile information.','A notification feed, responsive layout, summary statistics, data tables, and bar charts.','Persistence of demo workspace state in the browser through local storage, plus JSON export and reset controls.']: bullet(x)
p('The project does not implement backend APIs, a database, real authentication, authorization on a server, payment-gateway integration, geolocation, file uploads, messaging, real-time synchronization, production deployment, or real third-party marketplace data.')

page(); head('EXISTING AND PROPOSED SYSTEM')
p('In a conventional fragmented workflow, companies publish requirements through disconnected channels, contractors submit informal quotations, workers have limited visibility of assignments, and attendance or expenses are maintained in separate records. Information must be manually reconciled across people and tools.')
p('The proposed OBRIX system centralizes the demonstration flow in one interface. A company can publish a project; a contractor can submit a bid; the company can accept one bid; the resulting project becomes active in the contractor workspace; workers can be assigned and marked present; earnings and payments are then reflected in relevant dashboards. The same principle applies to equipment rental requests and lifecycle statuses.')

page(); head('FUNCTIONAL REQUIREMENTS')
reqs=[('FR1','Display public construction marketplace listings with search, city/category filters, and sorting.'),('FR2','Allow users to select a demo role and create a locally stored demo profile.'),('FR3','Allow companies to create, edit, view, and complete eligible projects.'),('FR4','Allow contractors to submit a single bid for an open project and track its status.'),('FR5','Allow companies to accept or reject pending bids and assign the chosen contractor to a project.'),('FR6','Allow contractors to assign workers to active projects and record present, absent, or half-day attendance.'),('FR7','Calculate worker earnings from attendance and agreed daily rates, and record simulated payments.'),('FR8','Allow contractors to record categorized site expenses and view expense charts.'),('FR9','Allow eligible users to request equipment rentals, and owners to approve, reject, or record returns.'),('FR10','Create and display notifications for selected workflow events.'),('FR11','Persist the workspace state in browser local storage and support JSON export/reset.')]
t=doc.add_table(rows=1, cols=2); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER; borders(t)
for i,x in enumerate(['ID','Requirement']): t.rows[0].cells[i].text=x; shade(t.rows[0].cells[i],'D9EAD3')
for a,b in reqs:
    c=t.add_row().cells; c[0].text=a; c[1].text=b
for row in t.rows:
    for c in row.cells:
        for para in c.paragraphs:
            for r in para.runs: font(r,10.5, row==t.rows[0])

page(); head('NON FUNCTIONAL REQUIREMENTS')
for x in ['Usability: concise role-specific navigation, labeled forms, feedback messages, empty states, and confirmation dialogs.','Responsiveness: CSS media queries adapt navigation, grids, forms, tables, and page spacing for smaller screens.','Accessibility: skip-to-content link, semantic navigation, ARIA labels where relevant, keyboard-aware dialogs, and reduced-motion support.','Performance: Vite bundling and lazy loading of primary public and workspace route components.','Data integrity within the demo: reducer actions validate required fields, numerical values, dates, role permissions, duplicate bids, payment limits, and rental state transitions.','Maintainability: functionality is separated across route, data, store, model, UI, public, workspace, finance, team, and rental components.']: bullet(x)

page(); head('SYSTEM MODULES')
modules=[('Public marketplace','Landing page, project/equipment/person listings, details, search, filters, sorting, and sign-in entry points.'),('Identity and roles','Demo login/signup plus role-aware routing for company, contractor, and worker workspaces.'),('Project and bid management','Project publication/editing, bid submission, bid decisions, contractor assignment, and progress/completion controls.'),('Workforce management','Worker directory, project assignment, attendance records, availability, skills, and reviews.'),('Finance','Payment recording, earnings calculation, wage balance display, expense records, filters, and charts.'),('Equipment rentals','Equipment browsing, rental request creation, owner decision, cancellation, and return recording.'),('Workspace services','Notifications, profile/settings screens, local persistence, export, reset, and UI feedback.')]
t=doc.add_table(rows=1,cols=2); t.style='Table Grid'; borders(t)
for i,x in enumerate(['Module','Responsibilities']): t.rows[0].cells[i].text=x; shade(t.rows[0].cells[i],'D9EAD3')
for a,b in modules: c=t.add_row().cells; c[0].text=a; c[1].text=b
for row in t.rows:
    for c in row.cells:
        for para in c.paragraphs:
            for r in para.runs: font(r,10.5,row==t.rows[0])

page(); head('METHODOLOGY')
p('The development follows an iterative frontend-development approach. Requirements are modeled as role-based user journeys, then expressed in reusable React components and reducer actions. Seed data provides a deterministic starting workspace. Each user action updates the central state, which causes connected screens to render the latest view.')
for x in ['Requirement analysis: define stakeholders, public discovery needs, and workspace responsibilities.','Information architecture and UI design: organize public pages and role-based navigation routes.','Component development: build reusable cards, forms, dialogs, tables, badges, charts, and layout components.','State modeling: define users, projects, bids, assignments, attendance, payments, expenses, equipment, rentals, notifications, and settings.','Validation and interaction logic: implement action-level checks and notification generation.','Testing and refinement: run unit/smoke/end-to-end tests and verify responsive user flows.']: bullet(x)

page(); head('TECHNOLOGY STACK')
stack=[('Frontend','React 18.3.1, React DOM 18.3.1, JSX'),('Build tool','Vite 6.4.1 with @vitejs/plugin-react 4.7.0'),('Routing','React Router DOM 6.28.0'),('Styling','CSS with responsive media queries; Tailwind CSS 3.4.17 is configured in the repository but the application styling is primarily authored in src/styles.css.'),('UI and icons','lucide-react 0.468.0'),('Charts','Recharts 2.15.0'),('State and storage','React context/store with reducer-style transitions, structured cloning, and browser localStorage'),('Testing','Node test runner, Playwright 1.51.1'),('Backend','Not used in the current project'),('Database','Not used in the current project'),('Deployment','Not configured in the current project; Vite preview is available for local preview.')]
t=doc.add_table(rows=1,cols=2); t.style='Table Grid'; borders(t)
for i,x in enumerate(['Layer','Technology / Status']): t.rows[0].cells[i].text=x; shade(t.rows[0].cells[i],'D9EAD3')
for a,b in stack: c=t.add_row().cells; c[0].text=a; c[1].text=b
for row in t.rows:
    for c in row.cells:
        for para in c.paragraphs:
            for r in para.runs: font(r,10.3,row==t.rows[0])

page(); head('HARDWARE AND SOFTWARE REQUIREMENTS')
p('Development and local execution requirements:')
for x in ['A modern desktop or laptop computer with an internet browser.','Node.js and npm compatible with the project dependencies.','A code editor such as Visual Studio Code (recommended, but not required by the source).','Operating system capable of running Node.js and a modern browser.','For end-to-end tests, Playwright browser dependencies as configured by the project environment.']: bullet(x)
p('No special server hardware, cloud account, database server, payment processor, or external API key is required for the present frontend demo.')

page(); head('SYSTEM DESIGN AND ARCHITECTURE')
p('OBRIX uses a client-side single-page architecture. The React entry point mounts the application. Route definitions render either public layouts or a protected role workspace. A shared store supplies the current user, seeded workspace data, and actions. Feature components request actions through the store; action transitions validate inputs and return a new state; the store persists the state in local storage.')
t=doc.add_table(rows=1,cols=3); t.style='Table Grid'; borders(t)
for i,x in enumerate(['Layer','Primary files','Purpose']): t.rows[0].cells[i].text=x; shade(t.rows[0].cells[i],'D9EAD3')
for a,b,cx in [('Application shell','main.jsx, App.jsx, routes.js','App mounting, routing, public layout, dashboard layout, lazy loading.'),('Shared presentation','UI.jsx, styles.css','Cards, forms, dialogs, tables, charts, visual assets, responsive styling.'),('Feature screens','Public.jsx, Workspace.jsx, Finance.jsx, Team.jsx, Rentals.jsx','Marketplace, role dashboards, finance, workforce, and rental experiences.'),('State and domain','Store.jsx, data.js, model.js','Seed data, local persistence, reducer actions, calculated earnings/payment helpers.')]:
    cells=t.add_row().cells; cells[0].text=a; cells[1].text=b; cells[2].text=cx
for row in t.rows:
    for c in row.cells:
        for para in c.paragraphs:
            for r in para.runs: font(r,9.5,row==t.rows[0])

page(); head('DATA AND STATE MANAGEMENT')
p('The application begins with a seeded in-memory model containing demo users, projects, bids, assignments, attendance, payments, expenses, equipment, rentals, notifications, reviews, and settings. The active session and all modified demo records are kept in a shared client-side state store and saved under the local-storage key obrix.workspace.v1.')
p('The state model is not a multi-user database. It is a simulation on one browser. Consequently, data changes do not synchronize across devices or browser profiles, and clearing browser storage can remove the saved demo state unless it has been exported as JSON.')

page(); head('TESTING STRATEGY')
p('The repository includes automated checks for core state behavior and representative routes. The package scripts provide Node-based tests, Playwright end-to-end tests, a Vite production build, and a local preview command.')
for x in ['Unit and model checks for state restoration and transition behavior.','Smoke coverage for public routes and role-specific workspace routes.','Playwright browser tests for application pages and interactions.','Manual responsive review at desktop and mobile breakpoints.','Validation-focused testing for project fields, bid limits, payment limits, role access, duplicate bids, attendance, and rental status transitions.']: bullet(x)
p('The synopsis records the intended testing approach; test outcomes depend on the local environment in which the project is run.')

page(); head('PROJECT TIMELINE')
t=doc.add_table(rows=1,cols=3); t.style='Table Grid'; borders(t)
for i,x in enumerate(['Phase','Planned activities','Suggested duration']): t.rows[0].cells[i].text=x; shade(t.rows[0].cells[i],'D9EAD3')
for a,b,cx in [('1. Analysis','Stakeholder study, requirements, user flows, and data model','Week 1'),('2. Design','Wireframes, route map, responsive layout, and component plan','Week 2'),('3. Marketplace and roles','Public pages, authentication demo, dashboards, navigation','Weeks 3-4'),('4. Workflows','Projects, bids, workforce, finance, and rentals','Weeks 5-6'),('5. QA and refinement','Validation, test execution, responsive/accessibility checks','Week 7'),('6. Documentation','Synopsis, screenshots, final demonstration, and submission','Week 8')]:
    cells=t.add_row().cells; cells[0].text=a; cells[1].text=b; cells[2].text=cx
for row in t.rows:
    for c in row.cells:
        for para in c.paragraphs:
            for r in para.runs: font(r,10.5,row==t.rows[0])

page(); head('EXPECTED OUTCOMES')
for x in ['A visually consistent, responsive construction marketplace and multi-role workspace.','A complete frontend demonstration of linked workflows from project discovery to bid decision, workforce tracking, and simulated financial records.','Clear separation of company, contractor, and worker responsibilities through role-based routes and navigation.','Improved visibility of mock project data, daily attendance, earnings, payments, expenses, equipment availability, and notifications.','A codebase that can serve as a basis for future backend integration and production hardening.']: bullet(x)

page(); head('LIMITATIONS AND FUTURE SCOPE')
p('Limitations of the present version include client-only storage, demo identities, seeded/mock data, simulated financial records, no backend API, no database, no live authentication, and no production deployment configuration. The local-store model is suitable for demonstrating connected frontend behavior but not for a live multi-user construction marketplace.')
p('Future work can introduce a secure backend and database, real user authentication and authorization, cloud deployment, payment-gateway integration, document uploads, real-time notifications, chat, rating moderation, location-aware matching, schedule/calendar management, analytics, audit logs, and role-based administrative controls.')

page(); head('CONCLUSION')
p('OBRIX demonstrates how a role-based frontend application can make construction collaboration easier to visualize. By combining project listings, bids, team management, attendance, simulated payments, expenses, equipment requests, and notifications, the prototype shows a coherent end-to-end experience for three connected user roles. The present implementation is deliberately limited to the frontend, but its component structure, state model, validations, and test setup provide a practical foundation for a future full-stack system.')

page(); head('REFERENCES')
for x in ['Project source code: OBRIX frontend repository, including package.json and src/ components.','React Documentation. https://react.dev/','Vite Documentation. https://vite.dev/','React Router Documentation. https://reactrouter.com/','Recharts Documentation. https://recharts.org/','Playwright Documentation. https://playwright.dev/','MDN Web Docs: Web Storage API. https://developer.mozilla.org/docs/Web/API/Web_Storage_API']:
    p(x)

for table in doc.tables:
    repeat_header(table.rows[0])
doc.save(OUT)
print(OUT)
